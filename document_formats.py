"""Format-aware, structure-preserving parsers for uploaded financial reports.

The PDF parser remains in :mod:`document_loader`.  This module handles office
and delimited formats without flattening table rows into unlabelled number
sequences.  It deliberately does not evaluate spreadsheet formulas or run
macros; only cached cell values are considered numeric evidence.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import zipfile
from bisect import bisect_right
from datetime import date, datetime, time
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

from document_loader import (
    DocumentChunk,
    DocumentProcessingError,
    ParsedBlock,
    ParsedDocument,
    ParsedPage,
    _bind_financial_table_row,
    _find_reliable_table_header,
    _has_ambiguous_repeated_period_header,
    _has_comparative_period_columns,
    _looks_like_financial_table_row,
    chunk_document,
    clean_text,
)
from retrieval.periods import extract_periods

logger = logging.getLogger(__name__)

SUPPORTED_DOCUMENT_EXTENSIONS = frozenset({".pdf", ".html", ".xlsx", ".docx", ".csv"})
MAX_ARCHIVE_MEMBERS = 20_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 1_000
MAX_SPREADSHEET_CELLS = 1_000_000
MAX_SPREADSHEET_COLUMNS = 512
MAX_SPREADSHEET_ROWS_PER_SHEET = 100_000
MAX_CELL_TEXT_CHARS = 4_000
MAX_MERGED_RANGES_PER_SHEET = 50_000
STRUCTURED_PARSER_VERSION = "financial-structured-formats-v3"


def validate_document_payload(filename: str, content: bytes) -> str:
    """Validate the supported format from both suffix and file signature.

    Returns the normalized format extension.  A suffix alone is never trusted
    because it controls which parser handles an uploaded file.
    """

    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
        raise DocumentProcessingError(f"Unsupported document format. Supported formats: {supported}")
    if not content:
        raise DocumentProcessingError("Document file is empty")

    if extension == ".pdf":
        if not content.startswith(b"%PDF-") or b"%%EOF" not in content[-1024:]:
            raise DocumentProcessingError("Uploaded content does not match a valid PDF file")
        try:
            import fitz

            with fitz.open(stream=content, filetype="pdf") as document:
                if document.needs_pass or document.is_encrypted:
                    raise DocumentProcessingError("Encrypted PDF documents are not supported")
                if document.page_count < 1:
                    raise DocumentProcessingError("Uploaded PDF must contain at least one page")
        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentProcessingError("Uploaded PDF could not be opened") from exc
        return extension

    if extension == ".html":
        _validate_html_payload(content)
        return extension

    if extension in {".xlsx", ".docx"}:
        expected_entry = "xl/workbook.xml" if extension == ".xlsx" else "word/document.xml"
        _validate_office_archive(content, expected_entry)
        return extension

    _validate_csv_payload(content)
    return extension


def load_structured_document_chunks(
    path: str | Path,
    *,
    chunk_size: int,
    overlap: int,
) -> list[DocumentChunk]:
    """Parse a supported non-PDF report and return provenance-bearing chunks."""

    source = Path(path)
    if not source.is_file():
        raise DocumentProcessingError(f"Document file not found: {source.name}")
    content = source.read_bytes()
    extension = validate_document_payload(source.name, content)
    if extension == ".html":
        document = _parse_html(source, content)
    elif extension == ".xlsx":
        document = _parse_xlsx(source)
    elif extension == ".docx":
        document = _parse_docx(source)
    elif extension == ".csv":
        document = _parse_csv(source, content)
    else:
        raise DocumentProcessingError(f"No structured parser is registered for {extension}")

    chunks = chunk_document(document, chunk_size=chunk_size, overlap=overlap)
    if not chunks:
        raise DocumentProcessingError(f"{extension.upper().lstrip('.')} produced no indexable content")
    return chunks


class _FinancialHtmlParser(HTMLParser):
    """Extract ordered filing text while keeping table rows together."""

    _BLOCK_TAGS = frozenset({"article", "div", "h1", "h2", "h3", "h4", "li", "p", "section"})
    _SKIP_TAGS = frozenset({"script", "style", "noscript", "svg"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, str]] = []
        self._buffer: list[str] = []
        self._cells: list[str] = []
        self._cell_buffer: list[str] = []
        self._skip_depth = 0
        self._in_row = False

    @staticmethod
    def _normalise(parts: list[str]) -> str:
        return re.sub(r"\s+", " ", " ".join(parts)).strip()

    def _flush_block(self) -> None:
        text = self._normalise(self._buffer)
        if text:
            self.blocks.append((text, "narrative"))
        self._buffer.clear()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.casefold()
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "tr":
            self._flush_block()
            self._in_row = True
            self._cells.clear()
            self._cell_buffer.clear()
        elif tag in {"td", "th"} and self._in_row:
            self._cell_buffer.clear()
        elif tag in self._BLOCK_TAGS:
            self._flush_block()

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if tag in self._SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag in {"td", "th"} and self._in_row:
            cell = self._normalise(self._cell_buffer)
            if cell:
                self._cells.append(cell)
            self._cell_buffer.clear()
        elif tag == "tr" and self._in_row:
            row = self._normalise(self._cells)
            if row:
                self.blocks.append((row, "table"))
            self._cells.clear()
            self._in_row = False
        elif tag in self._BLOCK_TAGS:
            self._flush_block()

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        if self._in_row:
            self._cell_buffer.append(data)
        else:
            self._buffer.append(data)

    def finish(self) -> list[tuple[str, str]]:
        self._flush_block()
        return self.blocks


def _parse_html(source: Path, content: bytes) -> ParsedDocument:
    try:
        text = content.decode("utf-8", errors="replace")
        parser = _FinancialHtmlParser()
        parser.feed(text)
        blocks = parser.finish()
    except Exception as exc:
        raise DocumentProcessingError("HTML financial report could not be parsed") from exc
    parsed_blocks = tuple(
        ParsedBlock(
            text=clean_text(block_text),
            page=0,
            section="SEC filing table" if content_type == "table" else "SEC filing",
            ocr_used=False,
            content_type=content_type,
            source_locator="HTML table row" if content_type == "table" else None,
            table_context="SEC filing table" if content_type == "table" else None,
        )
        for block_text, content_type in blocks
        if len(clean_text(block_text)) >= 2
    )
    if not parsed_blocks:
        raise DocumentProcessingError("HTML financial report contains no readable text")
    return ParsedDocument(
        filename=source.name,
        pages=(ParsedPage(number=0, blocks=parsed_blocks, ocr_used=False),),
        parser_version="sec-html-v1",
        content_type="text/html",
    )


def _validate_html_payload(content: bytes) -> None:
    if len(content) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise DocumentProcessingError("HTML document exceeds the safe parsing limit")
    sample = content[:4096].lower()
    if b"<html" not in sample and b"<!doctype" not in sample:
        raise DocumentProcessingError("Uploaded content does not match an HTML document")
    parser = _FinancialHtmlParser()
    parser.feed(content.decode("utf-8", errors="replace"))
    if sum(len(text) for text, _ in parser.finish()) < 100:
        raise DocumentProcessingError("HTML document contains no meaningful report text")


def _validate_office_archive(content: bytes, expected_entry: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_MEMBERS:
                raise DocumentProcessingError("Office document contains too many archive entries")
            total_uncompressed = sum(entry.file_size for entry in entries)
            if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise DocumentProcessingError("Office document expands beyond the safe parsing limit")
            for entry in entries:
                if entry.file_size and entry.compress_size == 0:
                    raise DocumentProcessingError("Office document contains an invalid compressed entry")
                if entry.compress_size and entry.file_size / entry.compress_size > MAX_ARCHIVE_COMPRESSION_RATIO:
                    raise DocumentProcessingError("Office document exceeds the safe compression ratio")
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or expected_entry not in names:
                raise DocumentProcessingError("Uploaded content does not match its Office file extension")
    except DocumentProcessingError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise DocumentProcessingError("Office document is not a readable ZIP-based file") from exc


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = content.decode(encoding)
            if text.strip():
                return text
        except UnicodeDecodeError:
            continue
    raise DocumentProcessingError("Text report must be non-empty UTF-8 or UTF-16")


def _validate_csv_payload(content: bytes) -> None:
    text = _decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = [
        row for row in csv.reader(io.StringIO(text), dialect)
        if any(value.strip() for value in row)
    ]
    if len(rows) < 2:
        raise DocumentProcessingError("CSV report must contain a header and at least one data row")
    # Financial exports commonly put a title, unit note, or issuer name on
    # their own first line before the actual column headings.  The parser
    # already searches the first five rows for a reliable period header, so
    # validation must not reject that supported layout before parsing starts.
    has_header_and_data = any(
        len(header) >= 2
        and any(len(data_row) >= 2 for data_row in rows[index + 1 :])
        for index, header in enumerate(rows[:5])
    )
    if not has_header_and_data:
        raise DocumentProcessingError("CSV report must contain at least two columns")


def _parse_csv(source: Path, content: bytes) -> ParsedDocument:
    text = _decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [[_clean_cell(value) for value in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2 or max((len(row) for row in rows), default=0) < 2:
        raise DocumentProcessingError("CSV report must contain a header and at least one data row")

    header_index = _find_reliable_table_header(rows[:5])
    header = rows[header_index] if header_index is not None else rows[0]
    header_rows = rows[: max(1, (header_index or 0) + 1)]
    header_context = "CSV headers: " + " | ".join(
        "; ".join(value for value in row if value) for row in header_rows
    )
    comparative_context = _has_comparative_period_columns(
        " ".join(value for row in rows[:5] for value in row)
    ) or _has_ambiguous_repeated_period_header(rows[:5])
    blocks: list[ParsedBlock] = []
    data_rows = rows[header_index + 1 :] if header_index is not None else rows[1:]
    first_data_row_number = (header_index or 0) + 2
    for row_number, row in enumerate(data_rows, start=first_data_row_number):
        bound = _bind_financial_table_row(header, row)
        rendered = f"Financial table row — {bound}" if bound else _render_delimited_row(header, row)
        if not rendered:
            continue
        content_type = "table"
        if bound is None and comparative_context and _looks_like_financial_table_row(" ".join(row)):
            content_type = "unverified_table"
        blocks.append(
            ParsedBlock(
                text=rendered,
                page=0,
                section="CSV table",
                ocr_used=False,
                table_context=header_context,
                source_locator=f"CSV row {row_number}",
                content_type=content_type,
            )
        )
    if not blocks:
        raise DocumentProcessingError("CSV report contains no non-empty data rows")
    return ParsedDocument(
        filename=source.name,
        pages=(ParsedPage(number=0, blocks=tuple(blocks), ocr_used=False),),
        parser_version="csv-structured-v2",
        content_type="text/csv",
    )


def _parse_xlsx(source: Path) -> ParsedDocument:
    try:
        from openpyxl import load_workbook

        formulas_book = load_workbook(source, read_only=True, data_only=False)
        values_book = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentProcessingError("XLSX workbook could not be opened") from exc

    pages: list[ParsedPage] = []
    seen_cells = 0
    try:
        merged_rows_by_sheet = _xlsx_merged_row_ranges(
            source,
            formulas_book.worksheets,
        )
        for sheet_index, formulas_sheet in enumerate(formulas_book.worksheets, start=1):
            values_sheet = values_book[formulas_sheet.title]
            merged_row_intervals = merged_rows_by_sheet.get(formulas_sheet.title, ())
            merged_interval_starts = tuple(start for start, _end in merged_row_intervals)

            def row_has_merged_cells(row_number: int) -> bool:
                interval_index = bisect_right(merged_interval_starts, row_number) - 1
                return (
                    interval_index >= 0
                    and merged_row_intervals[interval_index][1] >= row_number
                )

            if formulas_sheet.max_column > MAX_SPREADSHEET_COLUMNS:
                raise DocumentProcessingError(
                    f"Worksheet {formulas_sheet.title!r} exceeds the {MAX_SPREADSHEET_COLUMNS}-column safety limit"
                )
            if formulas_sheet.max_row > MAX_SPREADSHEET_ROWS_PER_SHEET:
                raise DocumentProcessingError(
                    f"Worksheet {formulas_sheet.title!r} exceeds the {MAX_SPREADSHEET_ROWS_PER_SHEET}-row safety limit"
                )

            blocks: list[ParsedBlock] = []
            prior_nonempty_rows: list[str] = []
            prior_cell_rows: list[list[str]] = []
            period_header: list[str] | None = None
            period_header_row = 0
            period_header_ambiguous_until = 0
            row_streams = zip(
                formulas_sheet.iter_rows(),
                values_sheet.iter_rows(),
                strict=False,
            )
            for row_number, (formula_cells, value_cells) in enumerate(row_streams, start=1):
                rendered_cells: list[str] = []
                cell_values: list[str] = []
                for formula_cell, value_cell in zip(formula_cells, value_cells, strict=False):
                    if formula_cell.value is not None or value_cell.value is not None:
                        seen_cells += 1
                        if seen_cells > MAX_SPREADSHEET_CELLS:
                            raise DocumentProcessingError(
                                f"Workbook exceeds the {MAX_SPREADSHEET_CELLS}-cell safety limit"
                            )
                    is_formula = (
                        isinstance(formula_cell.value, str)
                        and formula_cell.value.startswith("=")
                    )
                    value = value_cell.value if is_formula else formula_cell.value
                    if is_formula:
                        value_text = (
                            f"{_format_cell_value(value)} [cached formula result]"
                            if value is not None
                            else "FORMULA_WITHOUT_CACHED_RESULT (not numeric evidence)"
                        )
                    else:
                        value_text = _format_cell_value(value)
                    cell_values.append(value_text)
                    if not value_text:
                        continue
                    number_format = formula_cell.number_format
                    format_note = (
                        f" [number format: {number_format}]"
                        if _has_financial_number_format(number_format)
                        else ""
                    )
                    rendered_cells.append(f"{formula_cell.column_letter}={value_text}{format_note}")
                if not rendered_cells:
                    continue

                has_reliable_period_header = _find_reliable_table_header([cell_values]) == 0
                has_period_labels = bool(extract_periods(" ".join(cell_values)))
                merged_current_row = row_has_merged_cells(row_number)
                if has_reliable_period_header and not merged_current_row:
                    period_header = cell_values
                    period_header_row = row_number
                    period_header_ambiguous_until = 0
                elif merged_current_row and (
                    has_reliable_period_header or has_period_labels
                ):
                    # A merged period header no longer proves a one-to-one
                    # column/value mapping. Keep it visible for retrieval and
                    # diagnostics, but prevent its nearby financial rows from
                    # becoming numeric evidence until an unmerged header is
                    # encountered.
                    period_header = None
                    period_header_row = row_number
                    period_header_ambiguous_until = row_number + 25

                current_row = f"Row {row_number}: " + " | ".join(rendered_cells)
                nearby_headers = "\n".join(prior_nonempty_rows[-4:])
                context = f"Worksheet: {formulas_sheet.title}\n"
                if nearby_headers:
                    context += f"Nearby title/header rows (original cell coordinates):\n{nearby_headers}\n"
                block_text = f"{context}{current_row}"
                content_type = "table"
                active_period_header = (
                    period_header
                    if period_header is not None
                    and 0 < row_number - period_header_row <= 25
                    else None
                )
                mapped_row = (
                    _bind_financial_table_row(active_period_header, cell_values)
                    if active_period_header is not None and row_number > period_header_row
                    else None
                )
                comparative_context = _has_comparative_period_columns(
                    " ".join(" ".join(row) for row in [*prior_cell_rows[-4:], cell_values])
                ) or _has_ambiguous_repeated_period_header([*prior_cell_rows[-4:], cell_values])
                looks_financial = _looks_like_financial_table_row(" ".join(cell_values))
                if mapped_row and not merged_current_row:
                    block_text = (
                        f"{context}Financial table row — {mapped_row}\n"
                        f"Source cells: {current_row}"
                    )
                elif looks_financial and (
                    merged_current_row
                    or row_number <= period_header_ambiguous_until
                    or comparative_context
                ):
                    content_type = "unverified_table"
                columns = [cell.split("=", 1)[0] for cell in rendered_cells]
                blocks.append(
                    ParsedBlock(
                        text=block_text,
                        page=0,
                        section=f"Worksheet: {formulas_sheet.title}",
                        ocr_used=False,
                        table_context=nearby_headers or None,
                        source_locator=(
                            f"Sheet {formulas_sheet.title!r}, row {row_number}, "
                            f"columns {columns[0]}–{columns[-1]}"
                        ),
                        content_type=content_type,
                    )
                )
                prior_nonempty_rows.append(current_row[:MAX_CELL_TEXT_CHARS])
                prior_cell_rows.append(cell_values)

            if blocks:
                pages.append(ParsedPage(number=sheet_index, blocks=tuple(blocks), ocr_used=False))
    finally:
        formulas_book.close()
        values_book.close()

    if not pages:
        raise DocumentProcessingError("XLSX workbook has no non-empty worksheet data")
    return ParsedDocument(
        filename=source.name,
        pages=tuple(pages),
        parser_version="xlsx-openpyxl-structured-v3",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _xlsx_merged_row_ranges(
    source: Path,
    worksheets: list[object],
) -> dict[str, tuple[tuple[int, int], ...]]:
    """Read merged-cell row spans without loading the workbook into memory.

    ``openpyxl``'s read-only worksheet intentionally omits merged-cell
    metadata. Financial sheets can use merged period headers or value cells,
    so the affected rows must be identified from the OOXML stream and
    quarantined instead of silently assigning a merged value to one period.
    """

    ranges_by_sheet: dict[str, tuple[tuple[int, int], ...]] = {}
    try:
        with zipfile.ZipFile(source) as archive:
            for worksheet in worksheets:
                worksheet_path = getattr(worksheet, "_worksheet_path", None)
                if (
                    not isinstance(worksheet_path, str)
                    or not worksheet_path.startswith("xl/worksheets/")
                    or ".." in Path(worksheet_path).parts
                ):
                    raise DocumentProcessingError(
                        "XLSX worksheet XML location could not be verified"
                    )
                try:
                    xml_stream = archive.open(worksheet_path)
                except KeyError as exc:
                    raise DocumentProcessingError(
                        "XLSX worksheet XML is missing"
                    ) from exc
                row_ranges: list[tuple[int, int]] = []
                with xml_stream:
                    element_stack = []
                    for event, element in ElementTree.iterparse(
                        xml_stream,
                        events=("start", "end"),
                    ):
                        if event == "start":
                            element_stack.append(element)
                            continue
                        if element.tag.rsplit("}", 1)[-1] == "mergeCell":
                            reference = element.attrib.get("ref", "")
                            row_numbers = re.findall(
                                r"\$?[A-Z]{1,3}\$?(\d+)",
                                reference,
                                flags=re.IGNORECASE,
                            )
                            if ":" not in reference or len(row_numbers) != 2:
                                raise DocumentProcessingError(
                                    "XLSX contains an invalid merged-cell range"
                                )
                            start_row, end_row = sorted(map(int, row_numbers))
                            if start_row < 1 or end_row > 1_048_576:
                                raise DocumentProcessingError(
                                    "XLSX contains an out-of-range merged-cell reference"
                                )
                            row_ranges.append((start_row, end_row))
                            if len(row_ranges) > MAX_MERGED_RANGES_PER_SHEET:
                                raise DocumentProcessingError(
                                    "XLSX worksheet contains too many merged-cell ranges"
                                )
                        element_stack.pop()
                        if element_stack:
                            element_stack[-1].remove(element)
                normalized_ranges: list[tuple[int, int]] = []
                for start_row, end_row in sorted(row_ranges):
                    if normalized_ranges and start_row <= normalized_ranges[-1][1] + 1:
                        previous_start, previous_end = normalized_ranges[-1]
                        normalized_ranges[-1] = (
                            previous_start,
                            max(previous_end, end_row),
                        )
                    else:
                        normalized_ranges.append((start_row, end_row))
                ranges_by_sheet[str(worksheet.title)] = tuple(normalized_ranges)
    except DocumentProcessingError:
        raise
    except (OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise DocumentProcessingError(
            "XLSX merged-cell structure could not be verified"
        ) from exc
    return ranges_by_sheet


def _parse_docx(source: Path) -> ParsedDocument:
    try:
        from docx import Document as load_docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        word_document = load_docx(source)
    except Exception as exc:
        raise DocumentProcessingError("DOCX document could not be opened") from exc

    blocks: list[ParsedBlock] = []
    current_section = "Document"
    current_unit_notes: list[str] = []
    paragraph_index = 0
    table_index = 0
    body = word_document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_index += 1
            paragraph = Paragraph(child, word_document)
            text = _clean_cell(paragraph.text)
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.casefold().startswith("heading") or style_name.casefold() == "title":
                current_section = text[:160]
                current_unit_notes = [text] if _is_financial_unit_note(text) else []
            elif _is_financial_unit_note(text):
                current_unit_notes = [*current_unit_notes[-2:], text]
            blocks.append(
                ParsedBlock(
                    text=f"{current_section}\n{text}",
                    page=0,
                    section=current_section,
                    ocr_used=False,
                    source_locator=f"DOCX paragraph {paragraph_index}",
                    content_type="narrative",
                )
            )
        elif child.tag.endswith("}tbl"):
            table_index += 1
            table = Table(child, word_document)
            table_rows = [
                (
                    [_clean_cell(cell.text) for cell in row.cells],
                    _docx_row_has_cell_spans(row),
                )
                for row in table.rows
            ]
            table_rows = [(row, has_spans) for row, has_spans in table_rows if any(row)]
            rows = [row for row, _has_spans in table_rows]
            row_has_spans = [has_spans for _row, has_spans in table_rows]
            if len(rows) < 2:
                continue
            header_index = _find_reliable_table_header(rows[:5])
            header = rows[header_index] if header_index is not None else rows[0]
            header_rows = rows[: max(1, (header_index or 0) + 1)]
            header_context = "DOCX table headers: " + " | ".join(
                "; ".join(value for value in row if value) for row in header_rows
            )
            header_has_spans = bool(
                header_index is not None and row_has_spans[header_index]
            )
            table_context = "\n".join([header_context, *current_unit_notes[-3:]])
            comparative_context = _has_comparative_period_columns(
                " ".join(value for row in rows[:5] for value in row)
            ) or _has_ambiguous_repeated_period_header(rows[:5])
            data_rows = rows[header_index + 1 :] if header_index is not None else rows[1:]
            for row_number, row in enumerate(data_rows, start=(header_index or 0) + 2):
                bound = _bind_financial_table_row(header, row)
                text = f"Financial table row — {bound}" if bound else _render_delimited_row(header, row)
                if not text:
                    continue
                content_type = "table"
                source_row_index = row_number - 1
                span_ambiguity = header_has_spans or (
                    source_row_index < len(row_has_spans)
                    and row_has_spans[source_row_index]
                )
                looks_financial = _looks_like_financial_table_row(" ".join(row))
                if (
                    (span_ambiguity and looks_financial)
                    or (bound is None and comparative_context and looks_financial)
                ):
                    content_type = "unverified_table"
                blocks.append(
                    ParsedBlock(
                        text=f"{current_section}\n{table_context}\n{text}",
                        page=0,
                        section=f"{current_section} / Table {table_index}",
                        ocr_used=False,
                        table_context=table_context,
                        source_locator=f"DOCX table {table_index}, row {row_number}",
                        content_type=content_type,
                    )
                )

    if not blocks:
        raise DocumentProcessingError("DOCX document has no extractable paragraphs or financial tables")
    return ParsedDocument(
        filename=source.name,
        pages=(ParsedPage(number=0, blocks=tuple(blocks), ocr_used=False),),
        parser_version="docx-python-docx-structured-v3",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _render_delimited_row(headers: list[str], row: list[str]) -> str:
    fields = []
    for index, value in enumerate(row):
        value = _clean_cell(value)
        if not value:
            continue
        header = _clean_cell(headers[index]) if index < len(headers) else ""
        column = header or f"column {index + 1}"
        fields.append(f"{column}: {value}")
    return "Financial table row — " + " | ".join(fields) if fields else ""


def _clean_cell(value: object) -> str:
    text = clean_text(str(value)) if value is not None else ""
    if len(text) > MAX_CELL_TEXT_CHARS:
        return text[:MAX_CELL_TEXT_CHARS] + " [cell text truncated]"
    return text


def _is_financial_unit_note(text: str) -> bool:
    return bool(
        re.search(
            r"\b(?:in|amounts?\s+(?:are\s+)?(?:stated\s+|reported\s+|presented\s+)?in|"
            r"expressed\s+in)\s+(?:(?:US\$?|USD)\s+)?"
            r"(?:thousands?|millions?|billions?|trillions?)\b",
            text,
            re.IGNORECASE,
        )
    )


def _docx_row_has_cell_spans(row: object) -> bool:
    """Return true when Word merged cells make this row's grid ambiguous.

    ``python-docx`` can expose merged cells as repeated values in ``row.cells``.
    Until the adapter preserves OOXML spans exactly, financial facts whose
    period header or value row uses a span are quarantined from the fact ledger.
    """

    for cell in getattr(getattr(row, "_tr", None), "tc_lst", ()):
        properties = getattr(cell, "tcPr", None)
        if properties is not None and (
            getattr(properties, "gridSpan", None) is not None
            or getattr(properties, "vMerge", None) is not None
        ):
            return True
    return False


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return format(value, ".15g")
    return _clean_cell(value)


def _has_financial_number_format(number_format: str) -> bool:
    lowered = number_format.casefold()
    return any(token in lowered for token in ("$", "€", "£", "¥", "%", "0,,", "0.0,,", "0,,,"))
